import argparse
import sqlite3
import subprocess
import sys
from datetime import date
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
DB_PATH = ROOT / "sales_database.db"
REPORTS_DIR = ROOT / "reports"
OUTPUT_PATH = REPORTS_DIR / "Retail_Sales_Analysis_Report.pptx"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
NAVY = RGBColor(28, 44, 68)
TEAL = RGBColor(29, 128, 134)
GREEN = RGBColor(42, 157, 143)
RED = RGBColor(214, 64, 69)
LIGHT_BG = RGBColor(245, 247, 250)
DARK_TEXT = RGBColor(31, 41, 55)
MUTED_TEXT = RGBColor(100, 116, 139)
WHITE = RGBColor(255, 255, 255)


def money(value):
    return f"INR {value:,.2f}"


def number(value):
    return f"{value:,.0f}"


def fetch_one(query):
    with sqlite3.connect(DB_PATH) as conn:
        conn.row_factory = sqlite3.Row
        return dict(conn.execute(query).fetchone())


def fetch_all(query):
    with sqlite3.connect(DB_PATH) as conn:
        conn.row_factory = sqlite3.Row
        return [dict(row) for row in conn.execute(query).fetchall()]


def run_selenium_smoke_test():
    test_file = ROOT / "tests" / "selenium_smoke_test.py"
    if not test_file.exists():
        return "NOT RUN", "tests/selenium_smoke_test.py was not found."

    result = subprocess.run(
        [sys.executable, str(test_file)],
        cwd=ROOT,
        capture_output=True,
        text=True,
        timeout=300,
    )
    output = "\n".join(part.strip() for part in [result.stdout, result.stderr] if part.strip())
    status = "PASS" if result.returncode == 0 else "FAIL"
    return status, output or "No output captured."


def add_background(slide, color=LIGHT_BG):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(slide, text, left, top, width, height, font_size=18, bold=False, color=DARK_TEXT):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.clear()
    p = frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    return box


def add_title(slide, title, subtitle=None):
    add_text(slide, title, Inches(0.55), Inches(0.35), Inches(8.8), Inches(0.55), 26, True, NAVY)
    if subtitle:
        add_text(slide, subtitle, Inches(0.58), Inches(0.88), Inches(8.8), Inches(0.35), 10, False, MUTED_TEXT)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(1.18), Inches(12.2), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = TEAL
    line.line.fill.background()


def add_card(slide, left, top, width, height, title, value, accent=TEAL):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = WHITE
    shape.line.color.rgb = RGBColor(219, 226, 235)

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.08), height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()

    add_text(slide, title, left + Inches(0.25), top + Inches(0.18), width - Inches(0.35), Inches(0.25), 10, True, MUTED_TEXT)
    add_text(slide, value, left + Inches(0.25), top + Inches(0.52), width - Inches(0.35), Inches(0.45), 20, True, NAVY)


def add_table(slide, rows, columns, left, top, width, height):
    table_shape = slide.shapes.add_table(len(rows) + 1, len(columns), left, top, width, height)
    table = table_shape.table

    for col_idx, column in enumerate(columns):
        cell = table.cell(0, col_idx)
        cell.text = column
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.color.rgb = WHITE
            paragraph.font.size = Pt(9)
            paragraph.font.bold = True

    for row_idx, row in enumerate(rows, start=1):
        for col_idx, column in enumerate(columns):
            cell = table.cell(row_idx, col_idx)
            cell.text = str(row[column])
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(9)
                paragraph.font.color.rgb = DARK_TEXT

    return table_shape


def add_image_fit(slide, image_path, left, top, max_width, max_height):
    image_path = Path(image_path)
    if not image_path.exists():
        add_text(slide, f"Missing image: {image_path.name}", left, top, max_width, Inches(0.4), 14, True, RED)
        return

    with Image.open(image_path) as image:
        width_px, height_px = image.size

    ratio = min(max_width / width_px, max_height / height_px)
    width = int(width_px * ratio)
    height = int(height_px * ratio)
    img_left = left + int((max_width - width) / 2)
    img_top = top + int((max_height - height) / 2)
    slide.shapes.add_picture(str(image_path), img_left, img_top, width=width, height=height)


def build_presentation(test_status, test_output):
    summary = fetch_one(
        """
        SELECT
            COUNT(*) AS total_orders,
            SUM(Revenue) AS total_revenue,
            SUM(Profit) AS total_profit,
            AVG(Revenue) AS avg_order_value
        FROM sales_data
        """
    )
    best_product = fetch_one(
        """
        SELECT Product_Name, Brand, SUM(Revenue) AS revenue
        FROM sales_data
        GROUP BY Product_Name, Brand
        ORDER BY revenue DESC
        LIMIT 1
        """
    )
    best_month = fetch_one(
        """
        SELECT SUBSTR(Date, 1, 7) AS month, SUM(Revenue) AS revenue
        FROM sales_data
        GROUP BY month
        ORDER BY revenue DESC
        LIMIT 1
        """
    )
    regions = fetch_all(
        """
        SELECT Region, ROUND(SUM(Revenue), 2) AS Revenue, ROUND(SUM(Profit), 2) AS Profit
        FROM sales_data
        GROUP BY Region
        ORDER BY Revenue DESC
        """
    )
    categories = fetch_all(
        """
        SELECT Category, ROUND(SUM(Revenue), 2) AS Revenue
        FROM sales_data
        GROUP BY Category
        ORDER BY Revenue DESC
        LIMIT 5
        """
    )

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, NAVY)
    add_text(slide, "Retail Sales Analysis", Inches(0.8), Inches(1.45), Inches(8.5), Inches(0.75), 36, True, WHITE)
    add_text(slide, "Executive dashboard report with model outputs and Selenium validation", Inches(0.84), Inches(2.2), Inches(8.5), Inches(0.4), 16, False, RGBColor(215, 226, 235))
    add_text(slide, f"Generated on {date.today().isoformat()}", Inches(0.86), Inches(6.45), Inches(4), Inches(0.3), 12, False, RGBColor(215, 226, 235))
    add_image_fit(slide, ROOT / "monthly_trend.png", Inches(7.0), Inches(1.05), Inches(5.4), Inches(4.9))

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    add_title(slide, "Executive Summary", "Live KPIs from sales_database.db")
    add_card(slide, Inches(0.7), Inches(1.55), Inches(3.0), Inches(1.25), "Total Orders", number(summary["total_orders"]), TEAL)
    add_card(slide, Inches(3.95), Inches(1.55), Inches(3.6), Inches(1.25), "Total Revenue", money(summary["total_revenue"]), GREEN)
    add_card(slide, Inches(7.8), Inches(1.55), Inches(3.6), Inches(1.25), "Total Profit", money(summary["total_profit"]), NAVY)
    add_card(slide, Inches(0.7), Inches(3.05), Inches(3.7), Inches(1.2), "Average Order Value", money(summary["avg_order_value"]), TEAL)
    add_card(slide, Inches(4.7), Inches(3.05), Inches(3.8), Inches(1.2), "Best Month", f"{best_month['month']} | {money(best_month['revenue'])}", GREEN)
    add_card(slide, Inches(8.8), Inches(3.05), Inches(3.55), Inches(1.2), "Top Product", best_product["Product_Name"], NAVY)
    add_text(slide, f"Top product brand: {best_product['Brand']} | Revenue: {money(best_product['revenue'])}", Inches(8.95), Inches(4.18), Inches(3.3), Inches(0.35), 9, False, MUTED_TEXT)
    add_table(slide, regions, ["Region", "Revenue", "Profit"], Inches(0.75), Inches(4.85), Inches(5.65), Inches(1.65))
    add_table(slide, categories, ["Category", "Revenue"], Inches(7.0), Inches(4.85), Inches(4.8), Inches(1.65))

    chart_slides = [
        ("Revenue Performance", "monthly_trend.png", "quarterly_analysis.png"),
        ("Product and Brand Performance", "product_revenue.png", "top_brands.png"),
        ("Channel, City, and Profit Insights", "channel_payment.png", "city_revenue.png"),
        ("Profitability Heatmap", "profit_heatmap.png", None),
        ("Forecast and Model Comparison", "forecast.png", "model_comparison.png"),
        ("Prediction Quality and Feature Importance", "actual_vs_predicted.png", "feature_importance.png"),
    ]

    for title, left_image, right_image in chart_slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_background(slide)
        add_title(slide, title)
        if right_image:
            add_image_fit(slide, ROOT / left_image, Inches(0.6), Inches(1.45), Inches(5.85), Inches(5.55))
            add_image_fit(slide, ROOT / right_image, Inches(6.85), Inches(1.45), Inches(5.85), Inches(5.55))
        else:
            add_image_fit(slide, ROOT / left_image, Inches(1.1), Inches(1.45), Inches(11.1), Inches(5.55))

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide)
    add_title(slide, "Automated Test Report", "Selenium smoke test for the Streamlit dashboard")
    status_color = GREEN if test_status == "PASS" else RED
    add_card(slide, Inches(0.75), Inches(1.6), Inches(3.1), Inches(1.2), "Selenium Result", test_status, status_color)
    add_text(slide, "Command", Inches(0.8), Inches(3.0), Inches(2.2), Inches(0.3), 12, True, NAVY)
    add_text(slide, "python tests\\selenium_smoke_test.py", Inches(0.8), Inches(3.35), Inches(6.4), Inches(0.35), 14, False, DARK_TEXT)
    add_text(slide, "Captured Output", Inches(0.8), Inches(4.05), Inches(2.4), Inches(0.3), 12, True, NAVY)
    output_box = add_text(slide, test_output[:1300], Inches(0.8), Inches(4.42), Inches(11.7), Inches(1.9), 11, False, DARK_TEXT)
    output_box.fill.solid()
    output_box.fill.fore_color.rgb = WHITE
    output_box.line.color.rgb = RGBColor(219, 226, 235)

    return prs


def main():
    parser = argparse.ArgumentParser(description="Generate a PowerPoint report for the retail sales analysis project.")
    parser.add_argument("--skip-tests", action="store_true", help="Do not rerun Selenium; mark test status as not run.")
    parser.add_argument("--output", default=str(OUTPUT_PATH), help="Output .pptx path.")
    args = parser.parse_args()

    REPORTS_DIR.mkdir(exist_ok=True)
    if args.skip_tests:
        test_status, test_output = "NOT RUN", "Selenium test was skipped."
    else:
        test_status, test_output = run_selenium_smoke_test()

    presentation = build_presentation(test_status, test_output)
    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output_path)
    print(f"Created PowerPoint report: {output_path}")
    print(f"Selenium test status included in report: {test_status}")


if __name__ == "__main__":
    main()
